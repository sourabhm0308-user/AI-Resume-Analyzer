import os
from fpdf import FPDF
import docx

os.makedirs('E:\\ai_resume_analyzer\\data\\sample_resumes', exist_ok=True)

resume_text = """
John Doe
Software Engineer
Email: john.doe@example.com | Phone: 555-0198 | LinkedIn: linkedin.com/in/johndoe

Summary
Passionate Software Engineer with 3 years of experience developing scalable web applications using Python, Django, and React. Strong background in database design and RESTful API development.

Experience
Software Engineer - Tech Solutions Inc.
Jan 2021 - Present
- Developed and maintained multiple web applications using React.js and Django.
- Improved database query performance by 30% through indexing and optimization.
- Collaborated with cross-functional teams to define, design, and ship new features.

Junior Developer - WebCorp
Jun 2019 - Dec 2020
- Assisted in the development of the company's main product using JavaScript and HTML/CSS.
- Wrote unit tests for frontend components, achieving 80% test coverage.

Education
Bachelor of Science in Computer Science
University of Technology - Graduated 2019

Skills
Languages: Python, JavaScript, HTML, CSS, SQL
Frameworks: Django, React, Node.js
Tools: Git, Docker, AWS
"""

# Create PDF
class PDF(FPDF):
    def header(self):
        pass
    def footer(self):
        pass

pdf = PDF()
pdf.add_page()
pdf.set_font("Arial", size=11)
for line in resume_text.split('\n'):
    pdf.cell(200, 7, txt=line, ln=1, align='L')
pdf.output("E:\\ai_resume_analyzer\\data\\sample_resumes\\sample_software_engineer.pdf")

# Create DOCX
doc = docx.Document()
for line in resume_text.split('\n'):
    doc.add_paragraph(line)
doc.save("E:\\ai_resume_analyzer\\data\\sample_resumes\\sample_software_engineer.docx")

print("Sample resumes created successfully.")
